"""
Formation de Cellules Industrielles & Optimiseur de Disposition
"""

from __future__ import annotations
import math, random, re
from collections import defaultdict
import matplotlib.pyplot as plt
import numpy as np
import pandas as pd
import seaborn as sns
import streamlit as st

st.set_page_config(page_title="Formation de Cellules Industrielles", layout="wide", page_icon="🏭", initial_sidebar_state="expanded")

st.markdown("""
<style>
@import url('https://fonts.googleapis.com/css2?family=Rajdhani:wght@400;500;600;700&family=Source+Sans+3:wght@300;400;600&display=swap');
:root{--primary:#1a3a5c;--accent:#e8721a;--light-bg:#f4f7fb;--card-bg:#ffffff;--border:#dce4ef;--text-main:#1a2b3c;--text-muted:#5a7184;--success:#2d9e6b;}
html,body,[class*="css"]{font-family:'Source Sans 3',sans-serif;color:var(--text-main);}
.app-header{background:linear-gradient(135deg,#1a3a5c 0%,#2e5f8a 50%,#1a3a5c 100%);padding:2rem 2.5rem 1.5rem;border-radius:12px;margin-bottom:2rem;display:flex;align-items:center;gap:1.5rem;border-bottom:4px solid var(--accent);position:relative;overflow:hidden;}
.app-header::before{content:'';position:absolute;top:-40px;right:-40px;width:200px;height:200px;background:rgba(232,114,26,0.08);border-radius:50%;}
.app-header-icon{font-size:3rem;filter:drop-shadow(0 2px 8px rgba(0,0,0,0.3));}
.app-header-text h1{font-family:'Rajdhani',sans-serif;font-weight:700;font-size:2rem;color:#ffffff;margin:0;letter-spacing:0.5px;}
.app-header-text p{color:rgba(255,255,255,0.75);margin:0.25rem 0 0;font-size:0.95rem;}
.step-header{display:flex;align-items:center;gap:1rem;padding:1rem 1.5rem;border-radius:10px;margin:1.5rem 0 1rem;}
.step-badge{font-family:'Rajdhani',sans-serif;font-weight:700;font-size:1rem;color:white;padding:0.3rem 0.8rem;border-radius:20px;white-space:nowrap;}
.step-title{font-family:'Rajdhani',sans-serif;font-weight:600;font-size:1.4rem;margin:0;}
.info-card{background:var(--card-bg);border:1px solid var(--border);border-radius:10px;padding:1.2rem 1.5rem;margin-bottom:1rem;box-shadow:0 2px 8px rgba(0,0,0,0.04);}
.kpi-grid{display:grid;grid-template-columns:repeat(auto-fit,minmax(130px,1fr));gap:0.8rem;margin:1rem 0;}
.kpi-box{background:var(--card-bg);border:1px solid var(--border);border-radius:8px;padding:0.9rem 1rem;text-align:center;box-shadow:0 1px 4px rgba(0,0,0,0.05);}
.kpi-value{font-family:'Rajdhani',sans-serif;font-size:1.8rem;font-weight:700;color:var(--primary);line-height:1.1;}
.kpi-label{font-size:0.75rem;color:var(--text-muted);text-transform:uppercase;letter-spacing:0.5px;margin-top:0.2rem;}
.ilot-card{border-radius:8px;padding:1rem 1.2rem;margin-bottom:0.8rem;border-left:5px solid;background:#f9fbff;}
.ilot-card h4{font-family:'Rajdhani',sans-serif;font-weight:600;font-size:1.05rem;margin:0 0 0.4rem;}
.ilot-id-panel{background:#0d1117;border-radius:10px;padding:1.5rem 2rem;font-family:'Courier New',monospace;font-size:0.88rem;color:#c9d1d9;margin:1rem 0;border:1px solid #30363d;line-height:1.8;}
.ilot-id-title{color:#58a6ff;font-weight:bold;font-size:0.95rem;text-align:center;border-bottom:1px solid #30363d;padding-bottom:0.8rem;margin-bottom:1rem;}
.ilot-ok{color:#3fb950;}
.ilot-warn{color:#f85149;}
.ilot-select{color:#ffa657;font-weight:bold;}
section[data-testid="stSidebar"]{background:linear-gradient(180deg,#1a3a5c 0%,#12293f 100%);}
section[data-testid="stSidebar"] *{color:rgba(255,255,255,0.9)!important;}
section[data-testid="stSidebar"] .stMarkdown p{color:rgba(255,255,255,0.75)!important;}
.sidebar-step{display:flex;align-items:center;gap:0.7rem;padding:0.6rem 0.8rem;border-radius:8px;margin-bottom:0.5rem;}
.sidebar-step.done{background:rgba(45,158,107,0.2);border:1px solid rgba(45,158,107,0.4);}
.sidebar-step.current{background:rgba(232,114,26,0.2);border:1px solid rgba(232,114,26,0.4);}
.sidebar-step.pending{background:rgba(255,255,255,0.05);border:1px solid rgba(255,255,255,0.1);}
.stButton>button{font-family:'Rajdhani',sans-serif;font-weight:600;font-size:1rem;border-radius:8px;transition:all 0.2s;}
.stButton>button[kind="primary"]{background:linear-gradient(135deg,#e8721a,#d4601a);border:none;color:white;}
.stButton>button[kind="primary"]:hover{transform:translateY(-1px);box-shadow:0 4px 12px rgba(232,114,26,0.35);}
.footer{text-align:center;padding:1.5rem;color:var(--text-muted);font-size:0.85rem;border-top:1px solid var(--border);margin-top:2rem;}
</style>
""", unsafe_allow_html=True)

MIN_CHAINON_MACHINES = 4
PALETTE = ["#1a3a5c","#e8721a","#2d9e6b","#9b59b6","#c0392b","#16a085","#d97706","#2980b9"]
PALETTE_LIGHT = ["#e8eef5","#fef0e6","#e8f7f2","#f5eef8","#fdecea","#e8f6f3","#fef9e7","#eaf4fb"]

# ── DATA PARSING ──────────────────────────────────────────────────────────

def parse_unique_machines(raw_seq: str) -> list:
    tokens = re.split(r"[;\-,\s]+", raw_seq.strip())
    seen, result = set(), []
    for t in tokens:
        t = t.strip().upper()
        if not t: continue
        if not t.startswith("M"): t = f"M{t}"
        if t not in seen:
            seen.add(t); result.append(t)
    return result

def load_vernicolor_excel(path) -> pd.DataFrame:
    from openpyxl import load_workbook
    wb = load_workbook(path, read_only=True)
    ws = wb.active
    rows = list(ws.iter_rows(values_only=True))
    gamme_start = None
    for i, row in enumerate(rows):
        if row and str(row[0]).strip().lower() == "produit":
            gamme_start = i + 1; break
    if gamme_start is None: gamme_start = 0
    records = []
    for row in rows[gamme_start:]:
        if not row or row[0] is None: continue
        product = str(row[0]).strip()
        if not product or product.lower() in ("produit","nan"): continue
        gamme = str(row[1]).strip() if len(row) > 1 and row[1] else ""
        if not gamme: continue
        qty = 1
        if len(row) > 2 and row[2] is not None:
            try: qty = int(row[2])
            except: qty = 1
        records.append({"Product": product, "Gamme": gamme, "Nombre de manutention": qty})
    return pd.DataFrame(records)

def build_incidence_matrix(gamme_df: pd.DataFrame) -> pd.DataFrame:
    product_machines = {}
    for _, row in gamme_df.iterrows():
        p = str(row.iloc[0]).strip()
        product_machines[p] = parse_unique_machines(str(row.iloc[1]))
    all_m = sorted({m for ms in product_machines.values() for m in ms}, key=lambda m: int(re.sub(r"\D","",m) or 0))
    matrix = {p: {m: (1 if m in product_machines[p] else 0) for m in all_m} for p in product_machines}
    return pd.DataFrame(matrix).T.fillna(0).astype(float)

# ── KING'S METHOD ─────────────────────────────────────────────────────────

def _king_sort_once(df):
    e = df.copy().astype(float); L, C = e.shape
    e["EDc"] = [sum(e.iloc[i,j]*(2**(C-j-1)) for j in range(C)) for i in range(L)]
    e.sort_values("EDc", ascending=False, inplace=True)
    ed_i = [sum(e.iloc[i,j]*(2**(L-i-1)) for i in range(L)) for j in range(C)] + [0]
    e.loc["EDi"] = ed_i
    sorted_cols = e.loc["EDi"].apply(pd.to_numeric,errors="coerce").sort_values(ascending=False).index
    return e[sorted_cols]

def apply_king_method(matrix_df):
    e = matrix_df.fillna(0).astype(float); L = len(e)
    e = _king_sort_once(e)
    for _ in range(100):
        prev_edi, prev_edc = e.iloc[L].tolist(), e["EDc"].tolist()
        e.drop("EDc",axis=1,inplace=True); e.drop(e.index[-1],axis=0,inplace=True)
        e = _king_sort_once(e)
        if e.iloc[L].tolist()==prev_edi and e["EDc"].tolist()==prev_edc: break
    final = e.copy()
    if "EDc" in final.columns: final.drop(columns=["EDc"],inplace=True)
    if "EDi" in final.index: final.drop(index="EDi",inplace=True)
    return final.astype(float)

# ── ISLAND DETECTION (dominant-assignment) ────────────────────────────────

def detect_cellules(gamme_df, king_df) -> list:
    machine_to_products = defaultdict(set)
    product_machines = {}
    for _, row in gamme_df.iterrows():
        p = str(row.iloc[0]).strip()
        mlist = parse_unique_machines(str(row.iloc[1]))
        product_machines[p] = mlist
        for m in mlist: machine_to_products[m].add(p)

    product_machine_count = {p: len(ms) for p, ms in product_machines.items()}

    # Order of appearance in dataset (for tie-breaking)
    product_order_idx = {p: i for i, p in enumerate(product_machines.keys())}

    machine_assignment = {}
    for m, prods in machine_to_products.items():
        if len(prods) == 1:
            machine_assignment[m] = list(prods)[0]
        else:
            # Assign to: 1) most machines (desc), 2) earliest in dataset (asc) — tie-break
            machine_assignment[m] = sorted(
                prods, key=lambda p: (-product_machine_count[p], product_order_idx.get(p, 999))
            )[0]

    island_machines = defaultdict(list)
    for m, p in machine_assignment.items(): island_machines[p].append(m)

    king_col_order = {m: i for i, m in enumerate(king_df.columns)}
    raw_islands = []
    for p, machines in island_machines.items():
        machines_sorted = sorted(machines, key=lambda m: king_col_order.get(m, 999))
        prods_in_island = sorted([prod for prod in product_machines if any(m in product_machines[prod] for m in machines)])
        raw_islands.append({"product": p, "machines": machines_sorted, "products": prods_in_island})
    raw_islands.sort(key=lambda x: king_col_order.get(x["machines"][0], 999))

    return [{"name": f"P{i+1}", "machines": isl["machines"], "products": isl["products"],
             "size": len(isl["machines"]), "usable": len(isl["machines"]) >= MIN_CHAINON_MACHINES}
            for i, isl in enumerate(raw_islands)]

# ── CHAÎNON METHOD ────────────────────────────────────────────────────────

def traffic_from_sequences(gamme_df, machines):
    machine_set = set(machines); lbl = {m: i for i, m in enumerate(machines)}
    T = np.zeros((len(machines), len(machines)), dtype=int)
    for _, row in gamme_df.iterrows():
        seq = [t for t in parse_unique_machines(str(row.iloc[1])) if t in machine_set]
        try: w = int(row.iloc[2])
        except: w = 1
        for k in range(len(seq)-1): T[lbl[seq[k]], lbl[seq[k+1]]] += w
    return T

def directed_edges_from_traffic(T, machines, threshold=1):
    edges, weights = [], {}
    for a in range(len(machines)):
        for b in range(len(machines)):
            if T[a,b] >= threshold:
                edges.append((machines[a], machines[b])); weights[(machines[a],machines[b])] = int(T[a,b])
    return edges, weights

def honeycomb_points(rows, cols, dx=1.2):
    dy = math.sqrt(3)/2*dx
    return {(r,c): (c*dx+(0.5*dx if r%2 else 0.0), -r*dy) for r in range(rows) for c in range(cols)}

def honeycomb_edges(rows, cols):
    E = set()
    for r in range(rows):
        for c in range(cols):
            if c+1<cols: E.add(((r,c),(r,c+1)))
            if r+1<rows:
                E.add(((r,c),(r+1,c)))
                if r%2==0 and c-1>=0: E.add(((r,c),(r+1,c-1)))
                elif r%2!=0 and c+1<cols: E.add(((r,c),(r+1,c+1)))
    return list(E)

def lattice_neighbors(cell, rows, cols):
    r,c = cell
    nbrs = [(r,c-1),(r,c+1)]
    for dr in [-1,1]:
        nr = r+dr
        if 0<=nr<rows:
            nbrs.append((nr,c))
            if r%2==0 and c-1>=0: nbrs.append((nr,c-1))
            elif r%2!=0 and c+1<cols: nbrs.append((nr,c+1))
    return [nb for nb in dict.fromkeys(nbrs) if 0<=nb[0]<rows and 0<=nb[1]<cols]

def bfs_cells(center, rows, cols, k):
    q,seen,out = [center],{center},[]
    while q and len(out)<k:
        cur=q.pop(0); out.append(cur)
        for nb in lattice_neighbors(cur,rows,cols):
            if nb not in seen: seen.add(nb); q.append(nb)
    return out

def _orient(a,b,c): return (b[0]-a[0])*(c[1]-a[1])-(b[1]-a[1])*(c[0]-a[0])
def _on_seg(a,b,c): return min(a[0],b[0])<=c[0]<=max(a[0],b[0]) and min(a[1],b[1])<=c[1]<=max(a[1],b[1])
def segs_intersect(p1,p2,q1,q2):
    o1,o2,o3,o4 = _orient(p1,p2,q1),_orient(p1,p2,q2),_orient(q1,q2,p1),_orient(q1,q2,p2)
    if (o1>0>o2 or o1<0<o2) and (o3>0>o4 or o3<0<o4): return True
    for oi,pi,s in [(o1,q1,(p1,p2)),(o2,q2,(p1,p2)),(o3,p1,(q1,q2)),(o4,p2,(q1,q2))]:
        if oi==0 and _on_seg(s[0],s[1],pi): return True
    return False

def layout_score(m2c, pts, edges, weights):
    pos = {m: pts[c] for m,c in m2c.items()}; vals = list(pos.values())
    cx,cy = sum(x for x,y in vals)/len(vals), sum(y for x,y in vals)/len(vals)
    dists = [math.hypot(x-cx,y-cy) for x,y in vals]
    length = sum(weights[(u,v)]*math.hypot(pos[u][0]-pos[v][0],pos[u][1]-pos[v][1]) for u,v in edges)
    segs = [(u,v,pos[u],pos[v]) for u,v in edges]
    crossings = sum(
        weights[(s[0],s[1])]+weights[(t[0],t[1])]
        for i,s in enumerate(segs) for t in segs[i+1:]
        if len({s[0],s[1],t[0],t[1]})==4 and segs_intersect(s[2],s[3],t[2],t[3])
    )
    pair = sum(math.hypot(vals[i][0]-vals[j][0],vals[i][1]-vals[j][1]) for i in range(len(vals)) for j in range(i+1,len(vals)))
    r = max(dists) if dists else 0
    return 140*crossings+0.6*length+20*sum(dists)+200*r**2+4*pair, crossings, length, sum(dists), r

def optimize_layout(machines, cells, pts, edges, weights, steps=80000, restarts=8, seed=0):
    rng = random.Random(seed); best,best_s,best_st = None,float("inf"),None
    T0,Te = 2.0,0.01
    for _ in range(restarts):
        c = cells[:]; rng.shuffle(c); cur = {machines[i]:c[i] for i in range(len(machines))}
        cs,*cst = layout_score(cur,pts,edges,weights); bl,bls,blst = cur.copy(),cs,cst
        for s in range(steps):
            t = T0*((Te/T0)**(s/max(1,steps-1))); a,b = rng.sample(machines,2); cur[a],cur[b]=cur[b],cur[a]
            ns,*nst = layout_score(cur,pts,edges,weights); d = ns-cs
            if d<=0 or rng.random()<math.exp(-d/max(t,1e-9)): cs=ns; (lambda:None)(); \
                (bls := ns) if ns<bls else None
            if ns<bls: bls,bl,blst=ns,cur.copy(),nst
            else: cur[a],cur[b]=cur[b],cur[a] if d>0 and rng.random()>=math.exp(-d/max(t,1e-9)) else (cur[a],cur[b])
        if bls<best_s: best_s,best,best_st = bls,bl,blst
    return best,best_s,best_st

def optimize_layout_fixed(machines, cells, pts, edges, weights, steps=80000, restarts=8, seed=0):
    rng = random.Random(seed); best,best_s,best_st = None,float("inf"),None; T0,Te = 2.0,0.01
    for _ in range(restarts):
        c = cells[:]; rng.shuffle(c); cur = {machines[i]:c[i] for i in range(len(machines))}
        cs,*cst = layout_score(cur,pts,edges,weights); bl,bls,blst = cur.copy(),cs,cst
        for s in range(steps):
            t = T0*((Te/T0)**(s/max(1,steps-1))); a,b = rng.sample(machines,2)
            cur[a],cur[b] = cur[b],cur[a]
            ns,*nst = layout_score(cur,pts,edges,weights); d = ns-cs
            accept = d<=0 or rng.random()<math.exp(-d/max(t,1e-9))
            if accept:
                cs = ns
                if ns<bls: bls,bl,blst = ns,cur.copy(),nst
            else:
                cur[a],cur[b] = cur[b],cur[a]
        if bls<best_s: best_s,best,best_st = bls,bl,blst
    return best,best_s,best_st

def _unit_perp(dx,dy):
    n=math.hypot(dx,dy); return (0.,0.) if n==0 else (-dy/n,dx/n)

def draw_chainon(m2c, pts, base_edges, edges, weights, title=""):
    fig,ax = plt.subplots(figsize=(14,8),facecolor="#f4f7fb"); ax.set_facecolor("#f4f7fb")
    for a,b in base_edges:
        x1,y1=pts[a]; x2,y2=pts[b]; ax.plot([x1,x2],[y1,y2],color="#c8d8e8",lw=0.8,zorder=1)
    for (r,c),(x,y) in pts.items():
        if (r,c) not in m2c.values(): ax.add_patch(plt.Circle((x,y),0.14,fc="#e8eef5",ec="#c0cfe0",lw=0.6,zorder=2))
    maxw = max(weights.values()) if weights else 1; edge_set = set(edges)
    NCOLS = ["#1a3a5c","#e8721a","#2d9e6b","#9b59b6","#c0392b","#16a085","#d97706","#2980b9"]
    for u,v in edges:
        x1,y1=pts[m2c[u]]; x2,y2=pts[m2c[v]]; dx,dy=x2-x1,y2-y1; dist=math.hypot(dx,dy)
        if dist==0: continue
        ux,uy=dx/dist,dy/dist; nr=0.3; start=(x1+ux*nr,y1+uy*nr); end=(x2-ux*nr,y2-uy*nr)
        offx,offy=0.,0.
        if (v,u) in edge_set:
            px,py=_unit_perp(dx,dy); sgn=1. if str(u)<str(v) else -1.; offx,offy=px*0.12*sgn,py*0.12*sgn
        w=weights[(u,v)]; lw=2.+4.5*(w/maxw); alpha=0.55+0.45*(w/maxw)
        ax.annotate("",xy=(end[0]+offx,end[1]+offy),xytext=(start[0]+offx,start[1]+offy),
            arrowprops=dict(arrowstyle="-|>",color="#e8721a",lw=lw,shrinkA=0,shrinkB=0,alpha=alpha),zorder=4)
        ax.text((start[0]+end[0])/2+offx,(start[1]+end[1])/2+offy,str(w),fontsize=8.5,ha="center",va="center",
            fontweight="bold",color="#1a3a5c",bbox=dict(boxstyle="round,pad=0.18",fc="white",ec="#dce4ef",lw=0.8),zorder=6)
    for i,(mid,cell) in enumerate(m2c.items()):
        x,y=pts[cell]; c=NCOLS[i%len(NCOLS)]
        ax.add_patch(plt.Circle((x,y),0.38,fc=c,ec="white",lw=0,zorder=6,alpha=0.12))
        ax.add_patch(plt.Circle((x,y),0.3,fc=c,ec="white",lw=2.5,zorder=7))
        ax.text(x,y,str(mid),ha="center",va="center",fontsize=9.5,fontweight="bold",color="white",zorder=8)
    ax.set_aspect("equal"); ax.axis("off"); ax.set_title(title,fontsize=12,fontweight="bold",color="#1a3a5c",pad=14)
    plt.tight_layout(pad=1.5); return fig

# ── SESSION STATE ─────────────────────────────────────────────────────────
for k,v in [("gamme_df",None),("incidence_df",None),("king_df",None),("cellules",None)]:
    if k not in st.session_state: st.session_state[k] = v

# ── HEADER ────────────────────────────────────────────────────────────────
st.markdown("""
<div class="app-header"><div class="app-header-icon">🏭</div>
<div class="app-header-text"><h1>Formation de Cellules Industrielles</h1>
<p>Optimisation de la disposition des ateliers · Méthode de King &amp; Méthode Chaînon</p></div></div>
""", unsafe_allow_html=True)

# ── SIDEBAR ───────────────────────────────────────────────────────────────
with st.sidebar:
    st.markdown("### 🏭 Navigation"); st.markdown("---")
    step_done = sum([st.session_state.gamme_df is not None, st.session_state.king_df is not None, st.session_state.cellules is not None])
    for i,(name,desc) in enumerate([("Données de gamme","Importer ou saisir les gammes"),("Méthode de King","Réordonner — détecter les îlots"),("Méthode Chaînon","Optimiser la disposition des flux")]):
        css = "done" if step_done>i else ("current" if i==step_done else "pending")
        icon = "✅" if step_done>i else ("▶️" if i==step_done else "⏳")
        st.markdown(f'<div class="sidebar-step {css}"><span style="font-size:1.1rem">{icon}</span><div><div style="font-weight:600;font-size:0.9rem">Étape {i+1} — {name}</div><div style="font-size:0.75rem;opacity:0.75">{desc}</div></div></div>',unsafe_allow_html=True)
    st.markdown("---"); st.caption("Méthode de King (îlots) + Méthode Chaînon (recuit simulé sur grille hexagonale).")
    if st.session_state.gamme_df is not None:
        st.markdown("---")
        if st.button("🔁 Réinitialiser tout",use_container_width=True):
            for k in ["gamme_df","incidence_df","king_df","cellules"]: st.session_state[k]=None
            st.rerun()

# ── ÉTAPE 0 ───────────────────────────────────────────────────────────────
st.markdown('<div class="step-header" style="background:#eef3fb"><span class="step-badge" style="background:#1a3a5c">Étape 0</span><span class="step-title" style="color:#1a3a5c">Saisie des Données de Gamme</span></div>',unsafe_allow_html=True)

if st.session_state.gamme_df is None:
    st.markdown('<div class="info-card">Fournissez le tableau de <strong>gamme opératoire</strong>. Le fichier Excel peut contenir un tableau de machines en haut et les gammes en dessous — la détection de section est automatique.</div>',unsafe_allow_html=True)
    col_ex,col_fmt = st.columns([2,1])
    with col_ex:
        st.markdown("**Format attendu :**")
        st.dataframe(pd.DataFrame({"Produit":["P1","P2"],"Gamme":["M1-M2-M3","M2-M4"],"Nb. manut.":[5,3]}),use_container_width=True,hide_index=True)
    with col_fmt:
        st.markdown("**Séparateurs :** `-` `;` `,` espace\n\n**Préfixe `M` :** optionnel\n\n**Doublons :** dédoublonnés auto.")
    st.markdown("---")
    mode = st.radio("Mode de saisie :",["📁 Importer un fichier Excel","✏️ Saisie manuelle"],horizontal=True)
    if mode == "📁 Importer un fichier Excel":
        uploaded = st.file_uploader("Glissez-déposez votre fichier Excel (.xlsx)",type=["xlsx","xls"])
        if uploaded:
            try:
                raw = load_vernicolor_excel(uploaded)
                if raw.empty: st.error("Aucune gamme détectée. Vérifiez le format.")
                else:
                    st.session_state.gamme_df = raw; st.success(f"✅ {len(raw)} produits chargés."); st.rerun()
            except Exception as e: st.error(f"Erreur : {e}")
    else:
        default = pd.DataFrame({"Product":["P1","P2","P3","P4"],"Gamme":["M1-M2-M3","M2-M4-M5","M1-M3-M5","M3-M4"],"Nombre de manutention":[5,3,4,2]})
        edited = st.data_editor(default,num_rows="dynamic",use_container_width=True,key="manual")
        if st.button("✅ Confirmer la saisie",type="primary"):
            edited = edited.dropna(subset=["Product","Gamme"])
            edited["Nombre de manutention"] = pd.to_numeric(edited["Nombre de manutention"],errors="coerce").fillna(1).astype(int)
            st.session_state.gamme_df = edited.reset_index(drop=True); st.success(f"✅ {len(edited)} produits."); st.rerun()
else:
    gdf = st.session_state.gamme_df
    all_m = {m for _,row in gdf.iterrows() for m in parse_unique_machines(str(row.iloc[1]))}
    st.markdown(f'<div class="kpi-grid"><div class="kpi-box"><div class="kpi-value">{len(gdf)}</div><div class="kpi-label">Produits</div></div><div class="kpi-box"><div class="kpi-value">{len(all_m)}</div><div class="kpi-label">Machines</div></div><div class="kpi-box"><div class="kpi-value">{gdf.iloc[:,2].sum()}</div><div class="kpi-label">Manut. totales</div></div></div>',unsafe_allow_html=True)
    st.success("✅ Données de gamme chargées avec succès.")
    with st.expander("📋 Aperçu du tableau de gamme",expanded=False):
        st.dataframe(gdf.rename(columns={"Product":"Produit","Nombre de manutention":"Nb. manutentions"}),use_container_width=True,hide_index=True)

# ── ÉTAPE 1 ───────────────────────────────────────────────────────────────
if st.session_state.gamme_df is not None:
    st.markdown('<div class="step-header" style="background:#fef3ea"><span class="step-badge" style="background:#e8721a">Étape 1</span><span class="step-title" style="color:#e8721a">Méthode de King — Formation des Îlots</span></div>',unsafe_allow_html=True)
    if st.session_state.king_df is None:
        st.markdown('<div class="info-card">La <strong>méthode de King</strong> réordonne la matrice d\'incidence produit × machine. Les machines partagées sont assignées à l\'îlot dominant (gamme la plus longue).</div>',unsafe_allow_html=True)
        if st.button("🔄 Appliquer la méthode de King",type="primary"):
            with st.spinner("Application en cours …"):
                try:
                    inc = build_incidence_matrix(st.session_state.gamme_df)
                    king = apply_king_method(inc)
                    cellules = detect_cellules(st.session_state.gamme_df, king)
                    st.session_state.incidence_df=inc; st.session_state.king_df=king; st.session_state.cellules=cellules
                    st.rerun()
                except Exception as e: st.error(f"Erreur : {e}")
    else:
        king = st.session_state.king_df; cellules = st.session_state.cellules
        n_usable = sum(1 for c in cellules if c["usable"])
        st.markdown(f'<div class="kpi-grid"><div class="kpi-box"><div class="kpi-value">{len(cellules)}</div><div class="kpi-label">Îlots détectés</div></div><div class="kpi-box"><div class="kpi-value">{len(king.columns)}</div><div class="kpi-label">Machines totales</div></div><div class="kpi-box"><div class="kpi-value" style="color:#2d9e6b">{n_usable}</div><div class="kpi-label">Utilisables Chaînon</div></div><div class="kpi-box"><div class="kpi-value">{MIN_CHAINON_MACHINES}</div><div class="kpi-label">Min. requis</div></div></div>',unsafe_allow_html=True)

        col_matrix, col_islands = st.columns([3,2])
        with col_matrix:
            st.markdown("**📊 Matrice d'incidence réordonnée**")
            n_cells = len(cellules)
            fig_hm,ax_hm = plt.subplots(figsize=(max(6,len(king.columns)*0.65),max(4,len(king.index)*0.55)),facecolor="#f4f7fb")
            ax_hm.set_facecolor("#f4f7fb")
            colour_mat = np.zeros_like(king.values,dtype=float)
            for ci,cell in enumerate(cellules):
                for j,col in enumerate(king.columns):
                    if col in cell["machines"]: colour_mat[:,j] = ci+1
            sns.heatmap(colour_mat,ax=ax_hm,cmap=sns.color_palette([PALETTE_LIGHT[i%len(PALETTE_LIGHT)] for i in range(n_cells+1)]),
                linewidths=0.8,linecolor="white",cbar=False,xticklabels=king.columns,yticklabels=king.index,
                annot=king.astype(int),fmt="d",annot_kws={"size":9,"color":"#1a3a5c","fontweight":"bold"},vmin=0,vmax=max(n_cells,1))
            for ci,cell in enumerate(cellules):
                idxs = [list(king.columns).index(m) for m in cell["machines"] if m in king.columns]
                if idxs:
                    ax_hm.text((min(idxs)+max(idxs))/2+0.5,-0.6,cell["name"],ha="center",va="bottom",fontsize=8,color=PALETTE[ci%len(PALETTE)],fontweight="bold")
            ax_hm.set_title("Matrice réordonnée — Îlots colorés",fontsize=11,fontweight="bold",color="#1a3a5c",pad=16)
            ax_hm.set_xlabel("Machines",fontsize=9,color="#5a7184"); ax_hm.set_ylabel("Produits",fontsize=9,color="#5a7184")
            ax_hm.tick_params(colors="#1a3a5c",labelsize=8); plt.tight_layout(pad=1.2); st.pyplot(fig_hm); plt.close(fig_hm)

        with col_islands:
            n_total = len(king.columns)
            auto_select = next((c for c in cellules if c["usable"]), None)
            rows_html = []
            for cell in cellules:
                m_str = ", ".join(cell["machines"])
                n = cell["size"]
                if cell["usable"]:
                    status = '<span class="ilot-ok">✅ UTILISABLE POUR CHAÎNONS</span>'
                else:
                    status = '<span class="ilot-warn">❌ TROP PETIT</span>'
                rows_html.append(f'▶ Îlot {cell["name"]} : {status}<br>&nbsp;&nbsp;Machines : {m_str} ({n} postes)')

            sel_html = ""
            if auto_select:
                m_list = str(auto_select["machines"])
                sel_html = (f'<br><span style="color:#f85149">{"!"*50}</span><br>'
                            f'<span class="ilot-select">SÉLECTION AUTOMATIQUE : Îlot {auto_select["name"]}</span><br>'
                            f"L'étude des chaînons se fera sur : {m_list}<br>"
                            f'<span style="color:#f85149">{"!"*50}</span>')

            panel = "<br><br>".join(rows_html) + sel_html
            st.markdown(
                f'<div class="ilot-id-panel"><div class="ilot-id-title">'
                f'{"="*50}<br>IDENTIFICATION AUTOMATIQUE DES ÎLOTS ({n_total} MACHINES)<br>{"="*50}'
                f'</div><br>{panel}</div>',
                unsafe_allow_html=True,
            )

            st.markdown("---")
            for idx,cell in enumerate(cellules):
                col_hex=PALETTE[idx%len(PALETTE)]; col_light=PALETTE_LIGHT[idx%len(PALETTE_LIGHT)]
                badge = (f'<span style="background:#2d9e6b;color:white;font-size:0.75rem;padding:2px 8px;border-radius:12px">✅ Utilisable</span>'
                         if cell["usable"] else
                         f'<span style="background:#c0392b;color:white;font-size:0.75rem;padding:2px 8px;border-radius:12px">❌ Trop petit</span>')
                st.markdown(
                    f'<div class="ilot-card" style="border-left-color:{col_hex};background:{col_light}">'
                    f'<h4 style="color:{col_hex}">🏭 Îlot {cell["name"]} &nbsp;{badge}</h4>'
                    f'<div style="font-size:0.85rem;color:#1a3a5c"><b>Machines ({cell["size"]}) :</b> {" · ".join(cell["machines"])}</div>'
                    f'<div style="font-size:0.82rem;color:#5a7184;margin-top:0.3rem"><b>Produits :</b> {", ".join(cell["products"])}</div>'
                    f'</div>',
                    unsafe_allow_html=True,
                )

# ── ÉTAPE 2 ───────────────────────────────────────────────────────────────
if st.session_state.cellules is not None:
    st.markdown('<div class="step-header" style="background:#eaf7f2"><span class="step-badge" style="background:#2d9e6b">Étape 2</span><span class="step-title" style="color:#2d9e6b">Méthode Chaînon — Optimisation des Flux</span></div>',unsafe_allow_html=True)
    cellules = st.session_state.cellules; gamme_df = st.session_state.gamme_df
    usable = [c for c in cellules if c["usable"]]
    if not usable:
        st.warning(f"Aucun îlot ne contient au moins {MIN_CHAINON_MACHINES} machines.")
    else:
        st.markdown('<div class="info-card">La <strong>méthode Chaînon</strong> optimise la disposition des machines sur une grille hexagonale (nid d\'abeille) par <strong>recuit simulé</strong>, minimisant les croisements et la distance pondérée totale.</div>',unsafe_allow_html=True)
        cell_labels = [f"Îlot {c['name']} — {', '.join(c['machines'])} ({c['size']} machines)" for c in usable]
        selected_label = st.selectbox("🔍 Sélectionner un îlot :", cell_labels)
        selected_cell = usable[cell_labels.index(selected_label)]
        cell_machines = selected_cell["machines"]; idx = cellules.index(selected_cell)
        col_hex = PALETTE[idx%len(PALETTE)]
        st.markdown(f'<div style="background:{PALETTE_LIGHT[idx%len(PALETTE_LIGHT)]};border:1px solid {col_hex};border-radius:8px;padding:0.7rem 1rem;margin:0.5rem 0"><span style="color:{col_hex};font-weight:600">Îlot {selected_cell["name"]} :</span> <span style="color:#1a3a5c">{" → ".join(cell_machines)}</span></div>',unsafe_allow_html=True)
        c1,c2 = st.columns(2)
        with c1: n_steps = st.slider("Nombre de pas :",20000,150000,80000,10000)
        with c2: n_restarts = st.slider("Redémarrages :",3,15,8)

        if st.button("🚀 Calculer la disposition Chaînon",type="primary"):
            with st.spinner(f"Recuit simulé en cours ({n_steps:,} pas × {n_restarts} redémarrages) …"):
                try:
                    machine_set = set(cell_machines)
                    filtered_rows = []
                    for _,row in gamme_df.iterrows():
                        rel = [m for m in parse_unique_machines(str(row.iloc[1])) if m in machine_set]
                        if len(rel)>=2:
                            try: qty = int(row.iloc[2])
                            except: qty = 1
                            filtered_rows.append([str(row.iloc[0]), "-".join(rel), qty])
                    if not filtered_rows:
                        st.warning("Aucun itinéraire produit ne passe par au moins 2 machines de cet îlot.")
                    else:
                        fdf = pd.DataFrame(filtered_rows,columns=["Product","Gamme","Nombre de manutention"])
                        T = traffic_from_sequences(fdf, cell_machines)
                        edges,weights = directed_edges_from_traffic(T, cell_machines)
                        if not edges:
                            st.warning("Aucun flux détecté entre les machines de cet îlot.")
                        else:
                            m_count = len(cell_machines)
                            rh = max(5,m_count+2); ch = max(6,m_count+3)
                            pts = honeycomb_points(rh,ch); base_edges_hc = honeycomb_edges(rh,ch)
                            center = (rh//2,ch//2); cands = bfs_cells(center,rh,ch,k=m_count)
                            best_layout,best_score,stats = optimize_layout_fixed(cell_machines,cands,pts,edges,weights,steps=n_steps,restarts=n_restarts,seed=42)
                            crossings,length,compact,radius = stats
                            st.markdown(f'<div class="kpi-grid"><div class="kpi-box"><div class="kpi-value" style="color:#2d9e6b">{best_score:.0f}</div><div class="kpi-label">Score optimal</div></div><div class="kpi-box"><div class="kpi-value" style="color:{"#c0392b" if crossings>0 else "#2d9e6b"}">{crossings:.0f}</div><div class="kpi-label">Croisements</div></div><div class="kpi-box"><div class="kpi-value">{length:.1f}</div><div class="kpi-label">Long. pondérée</div></div><div class="kpi-box"><div class="kpi-value">{radius:.2f}</div><div class="kpi-label">Rayon max</div></div></div>',unsafe_allow_html=True)
                            title = f"Îlot {selected_cell['name']} — Disposition optimisée  (croisements : {crossings:.0f}  |  rayon : {radius:.2f})"
                            fig = draw_chainon(best_layout,pts,base_edges_hc,edges,weights,title=title)
                            st.pyplot(fig); plt.close(fig)
                            ct,cf = st.columns(2)
                            with ct:
                                st.markdown("**📋 Matrice de trafic dirigé**")
                                tdf = pd.DataFrame(T,index=cell_machines,columns=cell_machines)
                                st.dataframe(tdf.style.background_gradient(cmap="Blues",axis=None),use_container_width=True)
                            with cf:
                                st.markdown("**🔗 Flux (par intensité)**")
                                st.dataframe(pd.DataFrame([{"De":u,"Vers":v,"Flux":weights[(u,v)]} for u,v in sorted(edges,key=lambda e:-weights[e])]),use_container_width=True,hide_index=True)
                except Exception as e: st.error(f"Erreur : {e}")

st.markdown('<div class="footer">🏭 Formation de Cellules Industrielles &nbsp;·&nbsp; Méthode de King + Méthode Chaînon<br><span style="font-size:0.8rem">Optimisation par recuit simulé · Grille hexagonale</span></div>',unsafe_allow_html=True)
